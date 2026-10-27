"""Importação do quadro de irmãos por planilha (CSV ou XLSX)."""
import csv
import io
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.db import transaction

from apps.core import perfis as P
from apps.core.utils import somente_digitos

from .models import CategoriaIrmao, Irmao

COLUNAS = ["nome_completo", "email", "grau", "situacao", "categoria", "cpf", "rg", "data_nascimento", "telefone",
           "whatsapp", "data_iniciacao", "data_promocao", "data_elevacao", "profissao", "formacao_academica",
           "empresa", "matricula"]
GRAUS = {"aprendiz": 1, "1": 1, "companheiro": 2, "2": 2, "mestre": 3, "3": 3}
SITUACOES = {v.lower(): k for k, v in Irmao.Situacao.choices} | {k: k for k, _ in Irmao.Situacao.choices}
MODELO_CSV = ";".join(COLUNAS) + "\nJosé da Silva;jose@exemplo.org;Mestre;ativo;Integral;123.456.789-09;12.345.678-9;" \
    "15/03/1975;(11) 91234-5678;(11) 91234-5678;10/05/2005;12/05/2006;20/06/2007;Engenheiro;Engenharia Civil;Construtora X;1234\n"


def _data(valor):
    valor = (valor or "").strip() if isinstance(valor, str) else valor
    if not valor:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(valor, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"data inválida: {valor!r} (use dd/mm/aaaa)")


def _linhas(arquivo):
    nome = arquivo.name.lower()
    dados = arquivo.read()
    if nome.endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(dados), read_only=True, data_only=True).active
        linhas = list(ws.iter_rows(values_only=True))
        cab = [str(c or "").strip().lower() for c in linhas[0]] if linhas else []
        return [dict(zip(cab, ["" if c is None else c for c in l])) for l in linhas[1:] if any(l)]
    texto = dados.decode("utf-8-sig", errors="replace")
    delim = ";" if texto.splitlines()[0].count(";") >= texto.splitlines()[0].count(",") else ","
    return [{k.strip().lower(): v for k, v in l.items()} for l in csv.DictReader(io.StringIO(texto), delimiter=delim)]


def importar_irmaos(arquivo, aplicar=False):
    """Valida (e opcionalmente grava) a planilha. Devolve (criados, atualizados, erros)."""
    if arquivo.size > 2 * 1024 * 1024:
        raise ValidationError("Planilha maior que 2 MB.")
    linhas = _linhas(arquivo)
    if len(linhas) > 2000:
        raise ValidationError("Limite de 2000 linhas por importação.")
    criados = atualizados = 0
    erros = []
    categorias = {c.nome.lower(): c for c in CategoriaIrmao.objects.all()}
    with transaction.atomic():
        for n, l in enumerate(linhas, start=2):
            try:
                nome = str(l.get("nome_completo", "")).strip()
                email = str(l.get("email", "")).strip().lower()
                if not nome or not email:
                    raise ValueError("nome_completo e email são obrigatórios")
                grau = GRAUS.get(str(l.get("grau", "")).strip().lower())
                if grau is None:
                    raise ValueError("grau deve ser Aprendiz, Companheiro ou Mestre")
                situacao = SITUACOES.get(str(l.get("situacao", "ativo")).strip().lower() or "ativo")
                if situacao is None:
                    raise ValueError("situação inválida")
                cat_nome = str(l.get("categoria", "")).strip().lower()
                categoria = categorias.get(cat_nome) if cat_nome else None
                if cat_nome and categoria is None:
                    raise ValueError(f"categoria inexistente: {l.get('categoria')}")
                campos = {
                    "nome_completo": nome, "grau": grau, "situacao": situacao, "categoria": categoria,
                    "cpf": str(l.get("cpf", "")).strip(), "rg": str(l.get("rg", "")).strip(),
                    "data_nascimento": _data(l.get("data_nascimento")), "telefone": str(l.get("telefone", "")).strip(),
                    "whatsapp": str(l.get("whatsapp", "")).strip(), "data_iniciacao": _data(l.get("data_iniciacao")),
                    "data_promocao": _data(l.get("data_promocao")), "data_elevacao": _data(l.get("data_elevacao")),
                    "profissao": str(l.get("profissao", "")).strip(),
                    "formacao_academica": str(l.get("formacao_academica", "")).strip(),
                    "empresa": str(l.get("empresa", "")).strip(), "matricula": str(l.get("matricula", "")).strip(),
                }
                existente = Irmao.objects.filter(email__iexact=email).first()
                if existente is None and campos["cpf"]:
                    existente = Irmao.objects.filter(cpf=campos["cpf"]).first()
                obj = existente or Irmao(email=email, perfis=[P.IRMAO])
                for k, v in campos.items():
                    if v not in ("", None) or not existente:
                        setattr(obj, k, v)
                obj.full_clean(exclude=["username", "password"])
                if not existente:
                    base = (email.split("@")[0] or "irmao")[:30]
                    username, i = base, 1
                    while Irmao.objects.filter(username=username).exists():
                        i += 1
                        username = f"{base}{i}"
                    obj.username = username
                    obj.set_unusable_password()
                if aplicar:
                    obj.save()
                criados += not existente
                atualizados += bool(existente)
            except (ValueError, ValidationError) as exc:
                msg = "; ".join(exc.messages) if isinstance(exc, ValidationError) else str(exc)
                erros.append((n, msg))
        if erros or not aplicar:
            transaction.set_rollback(True)
    return criados, atualizados, erros
